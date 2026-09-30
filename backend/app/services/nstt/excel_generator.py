"""
NSTT Excel Generator Service
============================
Generates NSTT_Output.xlsx workbook with two comprehensive sheets:
  1. Dashboard — Multi-level hierarchical report matching SampleOutput.xlsx structure,
                 including Summary Matrix, Failure Analysis table, and Wrong NSTT Exceptions.
  2. RAW       — Final Enriched Dataset containing all Namo columns + 5 Remedy columns.

Uses openpyxl with high-fidelity formatting, distinct section headers, borders, and auto-widths.
"""

from __future__ import annotations

import io
import logging
from typing import Any, Dict, List, Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.nstt.aggregator import build_nstt_aggregation

logger = logging.getLogger(__name__)


def generate_nstt_excel_workbook(
    master_records: List[Dict[str, Any]],
    aggregated_result: Optional[Dict[str, Any]] = None,
    stats: Optional[Dict[str, Any]] = None,
) -> io.BytesIO:
    """
    Constructs the complete NSTT_Output.xlsx in-memory buffer.

    Args:
        master_records: List of classified / enriched record dicts.
        aggregated_result: Optional precomputed hierarchical aggregation dict.
        stats: Optional stats dict from file ingestion.

    Returns:
        io.BytesIO positioned at 0.
    """
    if aggregated_result is None:
        aggregated_result = build_nstt_aggregation(master_records)

    wb = openpyxl.Workbook()

    # Style definitions
    font_family = "Calibri"
    
    title_font = Font(name=font_family, size=14, bold=True, color="FFFFFF")
    section_font = Font(name=font_family, size=11, bold=True, color="1E293B")
    header_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    sub_header_font = Font(name=font_family, size=9, bold=True, color="334155")
    bold_font = Font(name=font_family, size=9, bold=True, color="0F172A")
    regular_font = Font(name=font_family, size=9, color="1E293B")
    muted_font = Font(name=font_family, size=8, italic=True, color="64748B")

    title_fill = PatternFill(start_color="157A8A", end_color="157A8A", fill_type="solid")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    sub_header_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    accent_fill = PatternFill(start_color="F0FDFA", end_color="F0FDFA", fill_type="solid")
    highlight_fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    warning_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    thick_border_side = Side(border_style="medium", color="1E293B")
    
    cell_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )
    
    header_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thick_border_side,
        bottom=thick_border_side,
    )

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # =========================================================================
    # SHEET 1: DASHBOARD
    # =========================================================================
    ws_dash = wb.active
    ws_dash.title = "Dashboard"
    ws_dash.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_dash.merge_cells("A1:M1")
    title_cell = ws_dash["A1"]
    title_cell.value = "NSTT ANALYTICS — HIERARCHICAL CLASSIFICATION REPORT"
    title_cell.font = title_font
    title_cell.fill = title_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_dash.row_dimensions[1].height = 36

    # Subtitle / Summary Cards
    ws_dash.cell(row=2, column=1, value="Report Generated").font = muted_font
    ws_dash.cell(row=2, column=2, value="Matched Records").font = muted_font
    ws_dash.cell(row=2, column=3, value="Total SR").font = muted_font
    ws_dash.cell(row=2, column=4, value="NSTT Count").font = muted_font
    ws_dash.cell(row=2, column=5, value="NSTT %").font = muted_font
    ws_dash.cell(row=2, column=6, value="Automation %").font = muted_font
    ws_dash.cell(row=2, column=7, value="Manual %").font = muted_font

    ws_dash.cell(row=3, column=1, value="NSTT Module").font = bold_font
    ws_dash.cell(row=3, column=2, value=stats.get("matched", len(master_records)) if stats else len(master_records)).font = bold_font
    ws_dash.cell(row=3, column=3, value=aggregated_result.get("total_sr", 0)).font = bold_font
    ws_dash.cell(row=3, column=4, value=aggregated_result.get("nstt_count", 0)).font = bold_font
    ws_dash.cell(row=3, column=5, value=f"{aggregated_result.get('nstt_percentage', 0)}%").font = bold_font
    ws_dash.cell(row=3, column=6, value=f"{aggregated_result.get('automation', {}).get('percentage', 0)}%").font = bold_font
    ws_dash.cell(row=3, column=7, value=f"{aggregated_result.get('manual', {}).get('percentage', 0)}%").font = bold_font

    for col in range(1, 8):
        ws_dash.cell(row=2, column=col).fill = accent_fill
        ws_dash.cell(row=3, column=col).fill = accent_fill
        ws_dash.cell(row=2, column=col).border = cell_border
        ws_dash.cell(row=3, column=col).border = cell_border
        ws_dash.cell(row=2, column=col).alignment = align_center
        ws_dash.cell(row=3, column=col).alignment = align_center

    # Section 1 Header: Summary Matrix
    ws_dash.cell(row=5, column=1, value="1. NSTT HIERARCHICAL MATRIX (SampleOutput Structure)").font = section_font

    headers_l1 = [
        ("Total SR", 1, 2),
        ("NSTT Count", 2, 2),
        ("Percentage NSTT Count", 3, 2),
        ("Capture Type", 4, 1),
        ("", 5, 1),
        ("Sub-Bifurcation 1", 6, 1),
        ("", 7, 1),
        ("Factory", 8, 1),
        ("", 9, 1),
        ("Allocation", 10, 1),
        ("", 11, 1),
        ("Impact Breakdown", 12, 1),
        ("", 13, 1),
    ]

    headers_l2 = [
        "", "", "",
        "Type", "Count (%)",
        "Category", "Count (%)",
        "Type", "Count (%)",
        "Type", "Count",
        "SA (0)", "NSA (1)",
    ]

    r_h1 = 6
    r_h2 = 7
    ws_dash.row_dimensions[r_h1].height = 24
    ws_dash.row_dimensions[r_h2].height = 20

    for col_idx, text in enumerate(headers_l2, start=1):
        c = ws_dash.cell(row=r_h2, column=col_idx, value=text)
        c.font = sub_header_font
        c.fill = sub_header_fill
        c.border = cell_border
        c.alignment = align_center

    # Merged header Level 1
    ws_dash.merge_cells("A6:A7")
    ws_dash["A6"].value = "Total SR"
    ws_dash["A6"].font = header_font
    ws_dash["A6"].fill = header_fill
    ws_dash["A6"].alignment = align_center

    ws_dash.merge_cells("B6:B7")
    ws_dash["B6"].value = "NSTT Count"
    ws_dash["B6"].font = header_font
    ws_dash["B6"].fill = header_fill
    ws_dash["B6"].alignment = align_center

    ws_dash.merge_cells("C6:C7")
    ws_dash["C6"].value = "Percentage NSTT Count"
    ws_dash["C6"].font = header_font
    ws_dash["C6"].fill = header_fill
    ws_dash["C6"].alignment = align_center

    ws_dash.merge_cells("D6:E6")
    ws_dash["D6"].value = "Capture Type"
    ws_dash["D6"].font = header_font
    ws_dash["D6"].fill = header_fill
    ws_dash["D6"].alignment = align_center

    ws_dash.merge_cells("F6:G6")
    ws_dash["F6"].value = "Sub-Bifurcation 1"
    ws_dash["F6"].font = header_font
    ws_dash["F6"].fill = header_fill
    ws_dash["F6"].alignment = align_center

    ws_dash.merge_cells("H6:I6")
    ws_dash["H6"].value = "Factory"
    ws_dash["H6"].font = header_font
    ws_dash["H6"].fill = header_fill
    ws_dash["H6"].alignment = align_center

    ws_dash.merge_cells("J6:K6")
    ws_dash["J6"].value = "Allocation"
    ws_dash["J6"].font = header_font
    ws_dash["J6"].fill = header_fill
    ws_dash["J6"].alignment = align_center

    ws_dash.merge_cells("L6:M6")
    ws_dash["L6"].value = "Impact Breakdown"
    ws_dash["L6"].font = header_font
    ws_dash["L6"].fill = header_fill
    ws_dash["L6"].alignment = align_center

    # Helper data extraction
    total_sr = aggregated_result.get("total_sr", 0)
    nstt_count = aggregated_result.get("nstt_count", 0)
    nstt_pct = aggregated_result.get("nstt_percentage", 0)
    
    auto = aggregated_result.get("automation", {})
    same = auto.get("same_nstt", {})
    same_im = same.get("im", {})
    same_non_im = same.get("non_im", {})
    diff = auto.get("different_nstt", {})
    res_auto = auto.get("resolved_nstt", {})
    
    man = aggregated_result.get("manual", {})
    man_res = man.get("resolved", {})
    man_before = man.get("before_sr_creation", {})
    man_after = man.get("after_sr_creation", {})

    matrix_rows = [
        # Row 8: Auto -> Same -> IM -> Auto
        [
            total_sr, nstt_count, f"{nstt_pct}%",
            "Automation", f"{auto.get('count', 0)} ({auto.get('percentage', 0)}%)",
            "Same NSTT", f"{same.get('count', 0)} ({same.get('percentage', 0)}%)",
            "IM", f"{same_im.get('count', 0)} ({same_im.get('percentage', 0)}%)",
            "Auto", same_im.get("auto", {}).get("count", 0),
            same_im.get("auto", {}).get("sa", 0), same_im.get("auto", {}).get("nsa", 0),
        ],
        # Row 9: Auto -> Same -> IM -> Manual
        [
            "", "", "",
            "", "",
            "", "",
            "", "",
            "Manual", same_im.get("manual", {}).get("count", 0),
            same_im.get("manual", {}).get("sa", 0), same_im.get("manual", {}).get("nsa", 0),
        ],
        # Row 10: Auto -> Same -> Non IM -> Auto
        [
            "", "", "",
            "", "",
            "", "",
            "Non IM", f"{same_non_im.get('count', 0)} ({same_non_im.get('percentage', 0)}%)",
            "Auto", same_non_im.get("auto", {}).get("count", 0),
            same_non_im.get("auto", {}).get("sa", 0), same_non_im.get("auto", {}).get("nsa", 0),
        ],
        # Row 11: Auto -> Same -> Non IM -> Manual
        [
            "", "", "",
            "", "",
            "", "",
            "", "",
            "Manual", same_non_im.get("manual", {}).get("count", 0),
            same_non_im.get("manual", {}).get("sa", 0), same_non_im.get("manual", {}).get("nsa", 0),
        ],
        # Row 12: Auto -> Different NSTT
        [
            "", "", "",
            "", "",
            "Different NSTT", f"{diff.get('count', 0)} ({diff.get('percentage', 0)}%)",
            f"SA auto NSA: {diff.get('incident_sa_auto_nsa', 0)} | NSA auto SA: {diff.get('incident_nsa_auto_sa', 0)}", "",
            f"Both NSA: {diff.get('both_nsa', 0)} | Both SA: {diff.get('both_sa', 0)}", "",
            "—", "—",
        ],
        # Row 13: Auto -> Resolved NSTT
        [
            "", "", "",
            "", "",
            "Resolved NSTT", f"{res_auto.get('count', 0)} ({res_auto.get('percentage', 0)}%)",
            f"IM Auto: {res_auto.get('im', {}).get('auto', 0)} | IM Man: {res_auto.get('im', {}).get('manual', 0)}", "",
            f"Non-IM Auto: {res_auto.get('non_im', {}).get('auto', 0)} | Non-IM Man: {res_auto.get('non_im', {}).get('manual', 0)}", "",
            "—", "—",
        ],
        # Row 14: Manual -> Resolved
        [
            "", "", "",
            "Manual", f"{man.get('count', 0)} ({man.get('percentage', 0)}%)",
            "Resolved", f"{man_res.get('count', 0)} ({man_res.get('percentage', 0)}%)",
            "—", "—",
            "—", "—",
            man_res.get("sa", 0), man_res.get("nsa", 0),
        ],
        # Row 15: Manual -> Before SR Creation
        [
            "", "", "",
            "", "",
            "Before SR Creation", f"{man_before.get('count', 0)} ({man_before.get('percentage', 0)}%)",
            "—", "—",
            "—", "—",
            man_before.get("sa", 0), man_before.get("nsa", 0),
        ],
        # Row 16: Manual -> After SR Creation
        [
            "", "", "",
            "", "",
            "After SR Creation", f"{man_after.get('count', 0)} ({man_after.get('percentage', 0)}%)",
            "—", "—",
            "—", "—",
            man_after.get("sa", 0), man_after.get("nsa", 0),
        ],

    ]

    start_r = 8
    for i, rdata in enumerate(matrix_rows):
        cur_row = start_r + i
        ws_dash.row_dimensions[cur_row].height = 19
        for col_idx, val in enumerate(rdata, start=1):
            c = ws_dash.cell(row=cur_row, column=col_idx, value=val)
            c.font = regular_font
            c.border = cell_border
            c.alignment = align_center

    # Apply cell merges on matrix
    ws_dash.merge_cells("A8:A16")
    ws_dash.merge_cells("B8:B16")
    ws_dash.merge_cells("C8:C16")
    ws_dash.merge_cells("D8:D13")
    ws_dash.merge_cells("E8:E13")
    ws_dash.merge_cells("F8:F11")
    ws_dash.merge_cells("G8:G11")
    ws_dash.merge_cells("H8:H9")
    ws_dash.merge_cells("I8:I9")
    ws_dash.merge_cells("H10:H11")
    ws_dash.merge_cells("I10:I11")
    
    ws_dash.merge_cells("D14:D16")
    ws_dash.merge_cells("E14:E16")

    # Section 2: Failure Analysis Table
    r_fail_title = 19
    ws_dash.cell(row=r_fail_title, column=1, value="2. MANUAL POPULATION FAILURE ANALYSIS").font = section_font

    r_fail_h = 20
    ws_dash.cell(row=r_fail_h, column=1, value="Failure Category").font = header_font
    ws_dash.cell(row=r_fail_h, column=1).fill = header_fill
    ws_dash.cell(row=r_fail_h, column=1).alignment = align_left
    ws_dash.cell(row=r_fail_h, column=1).border = cell_border

    ws_dash.cell(row=r_fail_h, column=2, value="Count").font = header_font
    ws_dash.cell(row=r_fail_h, column=2).fill = header_fill
    ws_dash.cell(row=r_fail_h, column=2).alignment = align_center
    ws_dash.cell(row=r_fail_h, column=2).border = cell_border

    ws_dash.cell(row=r_fail_h, column=3, value="% of Failures").font = header_font
    ws_dash.cell(row=r_fail_h, column=3).fill = header_fill
    ws_dash.cell(row=r_fail_h, column=3).alignment = align_center
    ws_dash.cell(row=r_fail_h, column=3).border = cell_border

    failures = aggregated_result.get("failures", {})
    fail_grand = failures.get("grand_total", 0)

    def _fail_pct(v):
        return f"{round((v / fail_grand) * 100, 2)}%" if fail_grand else "0.0%"

    fail_rows = [
        ("NSTT Number not found in Remedy", failures.get("nstt_not_found", 0), _fail_pct(failures.get("nstt_not_found", 0))),
        ("Ring Failure", failures.get("ring_failure", 0), _fail_pct(failures.get("ring_failure", 0))),
        ("Section Failure", failures.get("section_failure", 0), _fail_pct(failures.get("section_failure", 0))),
        ("Unstitched", failures.get("unstitched", 0), _fail_pct(failures.get("unstitched", 0))),
    ]

    for i, (cat, cnt, pct) in enumerate(fail_rows, start=21):
        ws_dash.cell(row=i, column=1, value=cat).font = regular_font
        ws_dash.cell(row=i, column=1).alignment = align_left
        ws_dash.cell(row=i, column=1).border = cell_border

        ws_dash.cell(row=i, column=2, value=cnt).font = regular_font
        ws_dash.cell(row=i, column=2).alignment = align_center
        ws_dash.cell(row=i, column=2).border = cell_border

        ws_dash.cell(row=i, column=3, value=pct).font = regular_font
        ws_dash.cell(row=i, column=3).alignment = align_center
        ws_dash.cell(row=i, column=3).border = cell_border

    r_grand = 25
    ws_dash.cell(row=r_grand, column=1, value="Grand Total").font = bold_font
    ws_dash.cell(row=r_grand, column=1).fill = sub_header_fill
    ws_dash.cell(row=r_grand, column=1).alignment = align_left
    ws_dash.cell(row=r_grand, column=1).border = cell_border

    ws_dash.cell(row=r_grand, column=2, value=fail_grand).font = bold_font
    ws_dash.cell(row=r_grand, column=2).fill = sub_header_fill
    ws_dash.cell(row=r_grand, column=2).alignment = align_center
    ws_dash.cell(row=r_grand, column=2).border = cell_border

    ws_dash.cell(row=r_grand, column=3, value="100.0%").font = bold_font
    ws_dash.cell(row=r_grand, column=3).fill = sub_header_fill
    ws_dash.cell(row=r_grand, column=3).alignment = align_center
    ws_dash.cell(row=r_grand, column=3).border = cell_border

    # Section 3: Wrong NSTT Exceptions
    r_ex_title = 27
    ws_dash.cell(row=r_ex_title, column=1, value="3. WRONG NSTT EXCEPTIONS (ANG in TXN Attribution)").font = section_font

    r_ex1 = 28
    ws_dash.cell(row=r_ex1, column=1, value="Wrong NSTT Attached by Automation").font = regular_font
    ws_dash.cell(row=r_ex1, column=1).border = cell_border
    ws_dash.cell(row=r_ex1, column=2, value=aggregated_result.get("wrong_nstt", {}).get("automation_ang_txn", 0)).font = bold_font
    ws_dash.cell(row=r_ex1, column=2).border = cell_border
    ws_dash.cell(row=r_ex1, column=2).alignment = align_center

    r_ex2 = 29
    ws_dash.cell(row=r_ex2, column=1, value="Wrong NSTT Attached by Engineer").font = regular_font
    ws_dash.cell(row=r_ex2, column=1).border = cell_border
    ws_dash.cell(row=r_ex2, column=2, value=aggregated_result.get("wrong_nstt", {}).get("engineer_ang_txn", 0)).font = bold_font
    ws_dash.cell(row=r_ex2, column=2).border = cell_border
    ws_dash.cell(row=r_ex2, column=2).alignment = align_center

    # Adjust column widths for Dashboard
    for col in ws_dash.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_dash.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # =========================================================================
    # SHEET 2: RAW ENRICHED DATASET
    # =========================================================================
    ws_raw = wb.create_sheet(title="RAW")
    ws_raw.views.sheetView[0].showGridLines = True

    if master_records:
        raw_headers = list(master_records[0].keys())
        ws_raw.row_dimensions[1].height = 24

        # Header Row
        for col_idx, h in enumerate(raw_headers, start=1):
            cell = ws_raw.cell(row=1, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = header_border
            cell.alignment = Alignment(horizontal="left", vertical="center")

        # Data Rows
        for row_idx, rec in enumerate(master_records, start=2):
            for col_idx, h in enumerate(raw_headers, start=1):
                val = rec.get(h)
                if val is None:
                    val = ""
                elif isinstance(val, bool):
                    val = "TRUE" if val else "FALSE"
                cell = ws_raw.cell(row=row_idx, column=col_idx, value=val)
                cell.font = regular_font
                cell.border = cell_border
                cell.alignment = Alignment(horizontal="left", vertical="center")

        # Auto-adjust widths on RAW sheet (sample first 100 rows for speed)
        for col in ws_raw.iter_cols(min_row=1, max_row=min(len(master_records) + 1, 100)):
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_raw.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    output_stream = io.BytesIO()
    wb.save(output_stream)
    output_stream.seek(0)
    return output_stream
