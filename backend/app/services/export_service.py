import csv
import io
from datetime import datetime
from typing import Any, Dict, List, Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def generate_excel_export(
    columns: List[Dict[str, str]],
    data: List[Dict[str, Any]],
    title: Optional[str] = None,
    sheet_name: str = "Export",
) -> io.BytesIO:
    """
    Generates a professionally styled Excel workbook in memory.

    :param columns: List of dicts with 'key' and 'label' (e.g. [{'key': 'period', 'label': 'Time Period'}])
    :param data: List of dict rows matching column keys
    :param title: Optional report title displayed in top header banner
    :param sheet_name: Name of the worksheet
    :return: BytesIO buffer positioned at start
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (sheet_name or "Export")[:31]  # Excel sheet names max 31 chars
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    title_font = Font(name="Calibri", size=14, bold=True, color="1E293B")
    subtitle_font = Font(name="Calibri", size=9, italic=True, color="64748B")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    
    even_row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    odd_row_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    data_font = Font(name="Calibri", size=10, color="0F172A")
    
    thin_border_side = Side(border_style="thin", color="E2E8F0")
    cell_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )
    
    header_border_side = Side(border_style="medium", color="0F172A")
    header_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=header_border_side,
        bottom=header_border_side,
    )

    current_row = 1

    # Optional Title Banner
    if title:
        ws.cell(row=current_row, column=1, value=title).font = title_font
        current_row += 1
        ws.cell(
            row=current_row,
            column=1,
            value=f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} • Total Records: {len(data)}",
        ).font = subtitle_font
        current_row += 2  # Leave one blank row

    # Column Headers
    header_row_idx = current_row
    for col_idx, col_def in enumerate(columns, start=1):
        cell = ws.cell(row=header_row_idx, column=col_idx, value=col_def.get("label", col_def.get("key", "")))
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = header_border

    ws.row_dimensions[header_row_idx].height = 26

    # Data Rows
    current_row += 1
    for row_idx, row_item in enumerate(data, start=current_row):
        fill = even_row_fill if (row_idx % 2 == 0) else odd_row_fill
        ws.row_dimensions[row_idx].height = 20

        for col_idx, col_def in enumerate(columns, start=1):
            key = col_def.get("key", "")
            raw_val = row_item.get(key)
            if raw_val is None:
                val = "—"
            elif isinstance(raw_val, (int, float, str, bool)):
                val = raw_val
            elif isinstance(raw_val, list):
                val = ", ".join(str(x) for x in raw_val)
            else:
                val = str(raw_val)

            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.fill = fill
            cell.border = cell_border
            
            # Numeric / Duration / Period alignment
            if isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto-adjust column widths
    for col_idx, col_def in enumerate(columns, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(col_def.get("label", "")))
        
        # Check first 200 data values for length
        for r in data[:200]:
            v = str(r.get(col_def.get("key", ""), "") or "")
            if len(v) > max_len:
                max_len = len(v)
                
        ws.column_dimensions[col_letter].width = max(12, min(max_len + 5, 55))

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generate_csv_export(
    columns: List[Dict[str, str]],
    data: List[Dict[str, Any]],
) -> io.BytesIO:
    """
    Generates an RFC-compliant CSV buffer with UTF-8-SIG encoding.

    :param columns: List of dicts with 'key' and 'label'
    :param data: List of dict rows matching column keys
    :return: BytesIO buffer containing encoded CSV
    """
    str_buffer = io.StringIO()
    fieldnames = [c.get("label", c.get("key", "")) for c in columns]
    keys = [c.get("key", "") for c in columns]

    writer = csv.writer(str_buffer, quoting=csv.QUOTE_MINIMAL)
    writer.writerow(fieldnames)

    for row_item in data:
        row_vals = []
        for key in keys:
            raw_val = row_item.get(key)
            if raw_val is None:
                val = ""
            elif isinstance(raw_val, list):
                val = ", ".join(str(x) for x in raw_val)
            else:
                val = str(raw_val)
            row_vals.append(val)
        writer.writerow(row_vals)

    # Use utf-8-sig so Excel on Windows opens CSV correctly with UTF-8 characters
    bytes_buffer = io.BytesIO(str_buffer.getvalue().encode("utf-8-sig"))
    bytes_buffer.seek(0)
    return bytes_buffer
