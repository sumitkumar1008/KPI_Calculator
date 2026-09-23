import io
import pytest
import openpyxl


def test_export_excel_success(client):
    payload = {
        "format": "xlsx",
        "filename": "test_kpi_report",
        "title": "Test KPI Report",
        "sheet_name": "KPIs",
        "columns": [
            {"key": "period", "label": "Time Period"},
            {"key": "MTTI", "label": "MTTI (avg)"},
            {"key": "MTTR", "label": "MTTR (avg)"},
        ],
        "data": [
            {"period": "2026-08", "MTTI": "00:10:00", "MTTR": "01:00:00"},
            {"period": "2026-09", "MTTI": "00:08:30", "MTTR": "00:45:00"},
        ],
    }

    response = client.post("/api/v1/kpi/export", json=payload)
    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "test_kpi_report.xlsx" in response.headers.get("Content-Disposition", "")

    # Validate that the returned bytes are a valid openpyxl workbook
    wb = openpyxl.load_workbook(io.BytesIO(response.data))
    assert "KPIs" in wb.sheetnames
    ws = wb["KPIs"]
    # Check title and content
    assert ws.cell(row=1, column=1).value == "Test KPI Report"


def test_export_csv_success(client):
    payload = {
        "format": "csv",
        "filename": "test_kpi_report",
        "columns": [
            {"key": "period", "label": "Time Period"},
            {"key": "MTTI", "label": "MTTI"},
        ],
        "data": [
            {"period": "2026-08", "MTTI": "00:10:00"},
            {"period": "2026-09", "MTTI": "00:08:30"},
        ],
    }

    response = client.post("/api/v1/kpi/export", json=payload)
    assert response.status_code == 200
    assert "text/csv" in response.mimetype
    assert "test_kpi_report.csv" in response.headers.get("Content-Disposition", "")
    
    text = response.data.decode("utf-8-sig")
    assert "Time Period,MTTI" in text
    assert "2026-08,00:10:00" in text


def test_export_invalid_format(client):
    payload = {
        "format": "pdf",
        "columns": [{"key": "a", "label": "A"}],
        "data": [{"a": 1}],
    }
    response = client.post("/api/v1/kpi/export", json=payload)
    assert response.status_code == 400
    assert "Invalid export format" in response.get_json()["error"]


def test_export_missing_data(client):
    payload = {
        "format": "xlsx",
        "columns": [{"key": "a", "label": "A"}],
    }
    response = client.post("/api/v1/kpi/export", json=payload)
    assert response.status_code == 400
    assert "Payload missing 'data' array" in response.get_json()["error"]


def test_export_empty_data_success(client):
    payload = {
        "format": "xlsx",
        "columns": [{"key": "period", "label": "Period"}],
        "data": [],
    }
    response = client.post("/api/v1/kpi/export", json=payload)
    assert response.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(response.data))
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "Period"


def test_export_from_cache_flow(client):
    from app.services.cache_service import save_to_cache

    sample_rows = [
        {
            "SRNUMBER": "SR-101",
            "SRCREATIONTIME": "2026-08-10T10:00:00",
            "AUTOMATION_RUN": "Y",
            "AUTOMATION_RCA_CONCLUSION": "Y",
            "MTTI": "00:10:00",
            "MTTA": "00:05:00",
            "MTTAck": "00:02:00",
            "MTTR": "01:00:00",
            "MTTr": "00:30:00",
            "MTTI_seconds": 600,
            "MTTA_seconds": 300,
            "MTTAck_seconds": 120,
            "MTTR_seconds": 3600,
            "MTTr_seconds": 1800,
        },
        {
            "SRNUMBER": "SR-102",
            "SRCREATIONTIME": "2026-08-11T12:00:00",
            "AUTOMATION_RUN": "N",
            "AUTOMATION_RCA_CONCLUSION": "N",
            "MTTI": "00:20:00",
            "MTTA": "00:10:00",
            "MTTAck": "00:04:00",
            "MTTR": "02:00:00",
            "MTTr": "01:00:00",
            "MTTI_seconds": 1200,
            "MTTA_seconds": 600,
            "MTTAck_seconds": 240,
            "MTTR_seconds": 7200,
            "MTTr_seconds": 3600,
        },
    ]

    uid = save_to_cache(sample_rows)

    # 1. Check cache status
    status_resp = client.get("/api/v1/kpi/cache")
    assert status_resp.status_code == 200
    status_data = status_resp.get_json()
    assert status_data["status"] == "cached"
    assert status_data["row_count"] == 2

    # 2. Test cached export (Unified KPI average - Excel)
    resp_avg_xlsx = client.get(f"/api/v1/kpi/export/cache?type=avg&period=monthly&format=xlsx&upload_id={uid}")
    assert resp_avg_xlsx.status_code == 200
    assert resp_avg_xlsx.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    # 3. Test cached export (Raw data - CSV)
    resp_raw_csv = client.get("/api/v1/kpi/export/cache?type=raw&format=csv")
    assert resp_raw_csv.status_code == 200
    assert "text/csv" in resp_raw_csv.mimetype
    csv_text = resp_raw_csv.data.decode("utf-8-sig")
    assert "SR-101" in csv_text
    assert "SR-102" in csv_text

    # 4. Test cached export (Automation Run - Excel)
    resp_auto_run = client.get("/api/v1/kpi/export/cache?type=automation_run&period=monthly&format=xlsx")
    assert resp_auto_run.status_code == 200

    # 5. Test cached export (Automation RCA - CSV)
    resp_rca = client.get("/api/v1/kpi/export/cache?type=automation_rca&period=monthly&format=csv")
    assert resp_rca.status_code == 200

