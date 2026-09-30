"""
Tests for NSTT Excel Export Generator and Endpoint
==================================================
Verifies:
  - Workbook generation produces valid .xlsx BytesIO
  - Exactly 2 sheets present: 'Dashboard' and 'RAW'
  - 'Dashboard' sheet includes title, metrics, matrix table, failures, and exceptions
  - 'RAW' sheet includes all columns and row records
  - API endpoint POST /api/v1/nstt/export streams valid Excel attachment
"""

import io
import openpyxl
import pytest
from app import create_app
from app.services.nstt.excel_generator import generate_nstt_excel_workbook
from app.services.nstt.cache_store import clear_nstt_cache, save_nstt_to_cache


@pytest.fixture
def app():
    clear_nstt_cache()
    app = create_app({"TESTING": True, "DEBUG": False})
    yield app
    clear_nstt_cache()


@pytest.fixture
def client(app):
    return app.test_client()


@pytest.fixture
def sample_master_records():
    return [
        {
            "INCIDENTID": "INC001",
            "ATTRIBUTEDTO": "TXN",
            "SRCREATIONTIME": "2026-01-01 10:00:00",
            "FACTORY": "IM",
            "AUTO_ALLOCATION": "yes",
            "AUTOMATION_UPDATE_NSTT_NUMBER": "INC001",
            "AUTOMATION_RCA_CONCLUSION_TEXT": "",
            "ASSIGNED_SUPPORT_COMPANY": "XYZ",
            "UP_TIME": "2026-01-01 12:00:00",
            "Submit Date": "2026-01-01 09:00:00",
            "INCIDENT_IMPACT": "SA",
            "DESCRIPTION": "Test description 1",
            "is_total_sr": True,
            "is_nstt": True,
            "capture_type": "Automation",
            "nstt_type": "Same NSTT",
            "factory_type": "IM",
            "allocation_type": "Auto",
            "impact_type": "SA",
            "timing_type": None,
            "diff_nstt_category": None,
            "exception_type": None,
            "failure_category": None,
        },
        {
            "INCIDENTID": "INC002",
            "ATTRIBUTEDTO": "TNL",
            "SRCREATIONTIME": "2026-01-01 10:00:00",
            "FACTORY": "Non IM",
            "AUTO_ALLOCATION": "no",
            "AUTOMATION_UPDATE_NSTT_NUMBER": None,
            "AUTOMATION_RCA_CONCLUSION_TEXT": "",
            "ASSIGNED_SUPPORT_COMPANY": "XYZ",
            "UP_TIME": "2026-01-01 12:00:00",
            "Submit Date": "2026-01-01 09:00:00",
            "INCIDENT_IMPACT": "NSA",
            "DESCRIPTION": "Section cut occurred",
            "is_total_sr": True,
            "is_nstt": True,
            "capture_type": "Manual",
            "nstt_type": None,
            "factory_type": None,
            "allocation_type": None,
            "impact_type": "NSA",
            "timing_type": "Before SR Creation",
            "diff_nstt_category": None,
            "exception_type": None,
            "failure_category": "Section Failure",
        },
    ]


def test_excel_generator_structure(sample_master_records):
    buf = generate_nstt_excel_workbook(master_records=sample_master_records)
    assert isinstance(buf, io.BytesIO)
    buf.seek(0)

    wb = openpyxl.load_workbook(buf)
    sheet_names = wb.sheetnames
    assert "Dashboard" in sheet_names
    assert "RAW" in sheet_names
    assert len(sheet_names) == 2

    # Check Dashboard content
    ws_dash = wb["Dashboard"]
    assert "NSTT ANALYTICS" in str(ws_dash["A1"].value)
    assert ws_dash["A8"].value == 2  # Total SR

    # Check RAW content
    ws_raw = wb["RAW"]
    # Header row
    headers = [cell.value for cell in ws_raw[1]]
    assert "INCIDENTID" in headers
    assert "ATTRIBUTEDTO" in headers
    assert "FACTORY" in headers
    assert "DESCRIPTION" in headers

    # 2 record rows + 1 header row = 3 rows
    assert ws_raw.max_row == 3


def test_export_endpoint_with_cached_upload_id(client, sample_master_records):
    master_response = {
        "records": sample_master_records,
        "total_records": len(sample_master_records),
    }
    upload_id = save_nstt_to_cache(master_response, stats={"matched": 2})

    # Process first to cache classified & aggregated data
    client.post("/api/v1/nstt/process", json={"upload_id": upload_id})

    # Trigger export
    resp = client.post("/api/v1/nstt/export", json={"upload_id": upload_id})
    assert resp.status_code == 200
    assert "spreadsheetml.sheet" in resp.headers.get("Content-Type", "")
    assert "attachment" in resp.headers.get("Content-Disposition", "")

    # Validate returned binary
    excel_file = io.BytesIO(resp.data)
    wb = openpyxl.load_workbook(excel_file)
    assert "Dashboard" in wb.sheetnames
    assert "RAW" in wb.sheetnames


def test_export_endpoint_inline_master_response(client, sample_master_records):
    master_response = {
        "records": sample_master_records,
        "total_records": len(sample_master_records),
    }

    resp = client.post("/api/v1/nstt/export", json={"master_response": master_response})
    assert resp.status_code == 200
    assert "spreadsheetml.sheet" in resp.headers.get("Content-Type", "")
    excel_file = io.BytesIO(resp.data)
    wb = openpyxl.load_workbook(excel_file)
    assert wb["RAW"].max_row == 3
